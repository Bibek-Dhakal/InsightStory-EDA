"""Build the 5-slide executive deck (Situation, Complication, Data Analysis x2, Recommendation)."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from .config import Settings
from .insights import Insights

DARK = (0x1F, 0x4E, 0x79)
GREY = (0x55, 0x5F, 0x6B)
FOOTER = "Synthetic demo data | InsightStory-EDA"

DECK_NAME = "InsightStory-EDA_Executive_Deck.pptx"
SLIDE_W = 13.333
SLIDE_H = 7.5

KPI_LABELS = [
    ("Revenue", "revenue"),
    ("Profit margin", "margin_pct"),
    ("Orders", "orders"),
    ("Active customers", "active_customers"),
]
SITUATION_CHART = "01_revenue_trend.png"
# (insights titles/notes key, chart image) for slides 2-4
CHART_SLIDES = [
    ("complication", "07_churn_trend.png"),
    ("rfm", "03_rfm_segments.png"),
    ("margin", "02_category_margin.png"),
]


def _add_text(slide, x, y, w, h, lines, size=16, bold=False, color=GREY):
    """Add a textbox; `lines` is a list of strings (one paragraph each)."""
    frame = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
    frame.word_wrap = True
    for i, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        run = paragraph.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = RGBColor(*color)
    return frame


def _new_slide(prs, title):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_text(slide, 0.6, 0.35, 12.1, 1.0, [title], size=28, bold=True, color=DARK)
    _add_text(slide, 0.6, 7.0, 8.0, 0.4, [FOOTER], size=10)
    return slide


def fmt_kpi(name: str, value: float) -> str:
    if name == "revenue" or name == "profit":
        return f"${value / 1_000_000:.2f}M" if value >= 1_000_000 else f"${value:,.0f}"
    if name == "margin_pct":
        return f"{value:.1f}%"
    return f"{value:,.0f}"


def build_deck(ins: Insights, chart_dir: Path, cfg: Settings) -> Path:
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)

    # Slide 1: Situation
    slide = _new_slide(prs, ins.titles["situation"])
    for i, (label, key) in enumerate(KPI_LABELS):
        x = 0.6 + i * 3.1
        _add_text(slide, x, 1.35, 2.9, 0.6, [fmt_kpi(key, ins.kpis[key])], size=30, bold=True, color=DARK)
        _add_text(slide, x, 1.95, 2.9, 0.4, [label], size=14)
    slide.shapes.add_picture(str(chart_dir / SITUATION_CHART), Inches(0.6), Inches(2.5), width=Inches(8.0))
    _add_text(slide, 9.0, 2.6, 3.8, 3.5, ins.notes["situation"], size=15)

    # Slides 2-4: Complication and Data Analysis
    for key, image in CHART_SLIDES:
        slide = _new_slide(prs, ins.titles[key])
        slide.shapes.add_picture(str(chart_dir / image), Inches(0.6), Inches(1.4), width=Inches(8.6))
        _add_text(slide, 9.5, 1.6, 3.4, 4.5, ins.notes[key], size=15)

    # Slide 5: Recommendation
    slide = _new_slide(prs, ins.titles["recommendation"])
    for i, rec in enumerate(ins.recommendations):
        y = 1.5 + i * 1.8
        _add_text(slide, 0.6, y, 12.1, 0.5, [f"{i + 1}. {rec['title']}"], size=20, bold=True, color=DARK)
        _add_text(slide, 0.9, y + 0.55, 11.8, 1.1, [f"Evidence: {rec['evidence']}"], size=14)

    cfg.output_dir.mkdir(parents=True, exist_ok=True)
    path = cfg.output_dir / DECK_NAME
    prs.save(str(path))
    return path
